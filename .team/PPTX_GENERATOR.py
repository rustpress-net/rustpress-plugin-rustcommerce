#!/usr/bin/env python3
"""
Amenthyx AI Teams — Shared PPTX Status Report Generator
Usage: python PPTX_GENERATOR.py <report_number> <project_name> [team_dir]
"""

import sys
import os

sys.stdout.reconfigure(encoding='utf-8')

from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from datetime import datetime


# Color scheme
DARK_BG = RGBColor(0x1A, 0x1A, 0x2E)
ACCENT = RGBColor(0x00, 0xD2, 0xFF)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
GREEN = RGBColor(0x00, 0xC8, 0x53)
YELLOW = RGBColor(0xFF, 0xC1, 0x07)
RED = RGBColor(0xFF, 0x17, 0x44)
GRAY = RGBColor(0x88, 0x88, 0x88)


def add_dark_slide(prs, title_text, subtitle_text=None):
    slide = prs.slides.add_slide(prs.slide_layouts[6])  # Blank layout
    bg = slide.background
    fill = bg.fill
    fill.solid()
    fill.fore_color.rgb = DARK_BG

    txBox = slide.shapes.add_textbox(Inches(0.5), Inches(0.3), Inches(12), Inches(1))
    tf = txBox.text_frame
    p = tf.paragraphs[0]
    p.text = title_text
    p.font.size = Pt(32)
    p.font.color.rgb = ACCENT
    p.font.bold = True

    if subtitle_text:
        p2 = tf.add_paragraph()
        p2.text = subtitle_text
        p2.font.size = Pt(16)
        p2.font.color.rgb = WHITE

    return slide


def add_text(slide, left, top, width, height, text, size=14, color=WHITE, bold=False):
    txBox = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
    tf = txBox.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = text
    p.font.size = Pt(size)
    p.font.color.rgb = color
    p.font.bold = bold
    return tf


def read_file_safe(path, max_chars=1500):
    if os.path.exists(path):
        with open(path, 'r', encoding='utf-8') as f:
            return f.read()[:max_chars]
    return None


def count_kanban(team_dir):
    kanban = read_file_safe(os.path.join(team_dir, "KANBAN.md"))
    if kanban:
        return {
            "done": kanban.count("[x]"),
            "in_progress": kanban.count("[~]"),
            "in_review": kanban.count("[?]"),
            "backlog": kanban.count("[ ]"),
        }
    return {"done": 0, "in_progress": 0, "in_review": 0, "backlog": 0}


def create_status_report(report_num, project_name, team_dir=".team"):
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    # Slide 1: Title
    slide1 = add_dark_slide(
        prs,
        f"Project Status Report #{report_num}",
        f"{project_name} | {datetime.now().strftime('%Y-%m-%d %H:%M')}"
    )
    add_text(slide1, 0.5, 2.5, 12, 1,
             "Amenthyx AI Teams — Virtual Engineering",
             size=20, color=GRAY)

    # Slide 2: Executive Summary
    slide2 = add_dark_slide(prs, "Executive Summary")
    kanban = count_kanban(team_dir)
    total = kanban["done"] + kanban["in_progress"] + kanban["in_review"] + kanban["backlog"]
    pct = int((kanban["done"] / total * 100)) if total > 0 else 0

    if pct >= 80:
        status_color = GREEN
        status_label = "ON TRACK"
    elif pct >= 50:
        status_color = YELLOW
        status_label = "IN PROGRESS"
    else:
        status_color = ACCENT
        status_label = "EARLY STAGE"

    add_text(slide2, 0.5, 1.8, 6, 0.6, f"Overall Status: {status_label}", size=20, color=status_color, bold=True)
    add_text(slide2, 0.5, 2.5, 12, 0.5, f"Completion: {pct}% ({kanban['done']}/{total} tasks done)", size=16, color=WHITE)
    add_text(slide2, 0.5, 3.2, 12, 0.5, f"In Progress: {kanban['in_progress']}  |  In Review: {kanban['in_review']}  |  Backlog: {kanban['backlog']}", size=14, color=GRAY)

    # Slide 3: Milestone Progress
    slide3 = add_dark_slide(prs, "Milestone Progress")
    milestones = read_file_safe(os.path.join(team_dir, "MILESTONES.md"))
    if milestones:
        add_text(slide3, 0.5, 1.8, 12, 5, milestones, size=11, color=WHITE)
    else:
        add_text(slide3, 0.5, 1.8, 12, 5, "Milestones pending — PM has not yet produced MILESTONES.md", size=14, color=YELLOW)

    # Slide 4: Kanban Snapshot
    slide4 = add_dark_slide(prs, "Kanban Snapshot")
    add_text(slide4, 0.5, 1.8, 3, 0.6, f"Backlog: {kanban['backlog']}", size=18, color=ACCENT)
    add_text(slide4, 3.5, 1.8, 3, 0.6, f"In Progress: {kanban['in_progress']}", size=18, color=YELLOW)
    add_text(slide4, 6.5, 1.8, 3, 0.6, f"In Review: {kanban['in_review']}", size=18, color=ACCENT)
    add_text(slide4, 9.5, 1.8, 3, 0.6, f"Done: {kanban['done']}", size=18, color=GREEN)

    kanban_content = read_file_safe(os.path.join(team_dir, "KANBAN.md"), 2000)
    if kanban_content:
        add_text(slide4, 0.5, 2.8, 12, 4, kanban_content, size=10, color=WHITE)

    # Slide 5: Blockers & Risks
    slide5 = add_dark_slide(prs, "Blockers & Risks")
    risks = read_file_safe(os.path.join(team_dir, "RISK_REGISTER.md"))
    if risks:
        add_text(slide5, 0.5, 1.8, 12, 5, risks, size=11, color=WHITE)
    else:
        add_text(slide5, 0.5, 1.8, 12, 5, "No blockers identified.", size=14, color=GREEN)

    # Slide 6: Team Activity
    slide6 = add_dark_slide(prs, "Team Activity")
    status_file = read_file_safe(os.path.join(team_dir, "TEAM_STATUS.md"))
    if status_file:
        add_text(slide6, 0.5, 1.8, 12, 5, status_file, size=10, color=WHITE)
    else:
        # Generic agent list
        agents = ["PM", "Backend", "Frontend", "Mobile", "DevOps", "Infra", "QA", "Release", "Marketing", "Legal"]
        y = 1.8
        for agent in agents:
            add_text(slide6, 0.5, y, 12, 0.4, f"  {agent}: Standby", size=12, color=GRAY)
            y += 0.45

    # Slide 7: Next Steps
    slide7 = add_dark_slide(prs, "Next Steps")
    add_text(slide7, 0.5, 1.8, 12, 5,
             "Planned for next reporting period:\n"
             "- Continue wave-based execution\n"
             "- Monitor quality gates\n"
             "- Update risk register\n"
             "- Coordinate cross-team dependencies\n"
             "- Generate next status report at T+6h",
             size=14, color=WHITE)

    # Save
    reports_dir = os.path.join(team_dir, "reports")
    os.makedirs(reports_dir, exist_ok=True)
    filepath = os.path.join(reports_dir, f"status_{report_num:03d}.pptx")
    prs.save(filepath)
    print(f"PPTX saved: {filepath}")
    return filepath


if __name__ == "__main__":
    report_num = int(sys.argv[1]) if len(sys.argv) > 1 else 1
    project_name = sys.argv[2] if len(sys.argv) > 2 else "Project"
    team_dir = sys.argv[3] if len(sys.argv) > 3 else ".team"
    create_status_report(report_num, project_name, team_dir)
